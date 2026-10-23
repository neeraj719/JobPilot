"""Text extraction from PDF and DOCX resumes.

Extraction is *best effort*: a resume that cannot be read (scanned image, corrupt
file, exotic encoding) still uploads successfully — the row is stored with
``parse_status=FAILED`` and a human-readable ``parse_error`` so the user can fix it.
"""
from __future__ import annotations

import io
import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

MIN_USEFUL_CHARS = 30


class ExtractionError(Exception):
    """Raised when a document's text cannot be read."""


#: Typographic characters that resumes are full of, mapped to plain text. Word and PDF
#: exporters emit em/en dashes, curly quotes, non-breaking and zero-width spaces, soft
#: hyphens and a zoo of bullet glyphs (Symbol/Wingdings private-use code points included);
#: the rules downstream (bullet splitting, keyword matching, number checks) and the UI want
#: one plain form. U+FFFD (a character some decoder already lost) carries nothing and is
#: dropped.
_CHAR_MAP: dict[int, str | None] = {
    **{ord(c): "-" for c in "\u2010\u2011\u2012\u2013\u2212\ufe58\ufe63\uff0d"},
    **{ord(c): " - " for c in "\u2014\u2015"},
    **{ord(c): "'" for c in "\u2018\u2019\u201a\u201b\u2032"},
    **{ord(c): '"' for c in "\u201c\u201d\u201e\u201f\u2033"},
    **{
        ord(c): "\u2022"
        for c in "\u25aa\u25cf\u25e6\u2023\u2043\u2219\u25a0\u25a1\u25c6\u2666\u27a2\u2794\u25ba\u2713\u2714"
        "\uf0b7\uf0a7\uf0d8\uf076\uf0fc"
    },
    **{ord(c): " " for c in "\u00a0\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u202f\u205f\u3000"},
    **{ord(c): None for c in "\u00ad\u200b\u200c\u200d\u2060\ufeff\ufffd"},
    ord("\u2026"): "...",
    # PDF WinAnsiEncoding also places the bullet at 0x7F (reportlab writes it there, so do
    # other generators); pypdf hands it back as DEL. Read it as the bullet it is.
    0x7F: "\u2022",
}
#: A bullet glyph at the start of a line becomes the plain marker "- ".
_LEADING_BULLET_RE = re.compile(r"^[ \t]*[\u2022\u00b7][ \t]*", re.MULTILINE)
#: A bullet glyph inside a line separates two items; keep it as " \u2022 " (``split_bullets`` splits on it).
_INLINE_BULLET_RE = re.compile(r"[ \t]*\u2022[ \t]*")
#: Control characters other than tab and newline (form feeds, vertical tabs, NULs from PDFs).
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


def normalize_text(text: str) -> str:
    """Plain-text form of typographic characters (see :data:`_CHAR_MAP`)."""
    text = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_RE.sub(" ", text.translate(_CHAR_MAP))
    text = _LEADING_BULLET_RE.sub("- ", text)
    return _INLINE_BULLET_RE.sub(" \u2022 ", text)


def _tidy(text: str) -> str:
    text = normalize_text(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" +\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ExtractionError("pypdf is not installed") from exc
    try:
        reader = PdfReader(io.BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - pypdf raises many error types
        raise ExtractionError(f"Could not read the PDF: {exc}") from exc
    return _tidy("\n".join(pages))


def _docx_paragraph_text(paragraph) -> str:  # noqa: ANN001 - python-docx Paragraph
    """Paragraph text; a Word list item (bullet numbering) keeps a "- " marker.

    Word draws list bullets from the numbering definition, so ``paragraph.text`` has no
    marker and a bullet would read like a heading line. The parser needs to tell them apart.
    """
    text = paragraph.text
    if not text.strip():
        return text
    try:
        ppr = paragraph._p.pPr  # noqa: SLF001 - python-docx has no public accessor
        is_list = ppr is not None and ppr.numPr is not None
        style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
        is_list = is_list or style_name.lower().startswith("list bullet")
    except Exception:  # noqa: BLE001 - never fail extraction over formatting details
        is_list = False
    if is_list and not re.match(r"^\s*(?:[-*•·]|\d+[.)])\s", text):
        return f"- {text}"
    return text


def extract_docx_text(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover - dependency is pinned
        raise ExtractionError("python-docx is not installed") from exc
    try:
        document = docx.Document(io.BytesIO(content))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many error types
        raise ExtractionError(f"Could not read the DOCX file: {exc}") from exc

    parts: list[str] = [_docx_paragraph_text(paragraph) for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)
    return _tidy("\n".join(parts))


def extract_text(content: bytes, ext: str) -> str:
    """Extract text from an in-memory document. Raises :class:`ExtractionError`."""
    kind = (ext or "").lower().lstrip(".")
    if kind == "pdf":
        text = extract_pdf_text(content)
    elif kind == "docx":
        text = extract_docx_text(content)
    else:
        raise ExtractionError(f"Unsupported document type: {ext!r}")
    if len(text) < MIN_USEFUL_CHARS:
        raise ExtractionError(
            "No readable text found in the document (a scanned image resume needs OCR, which is not supported)."
        )
    return text


def extract_text_from_path(path: str | Path) -> str:
    file_path = Path(str(path))
    if not file_path.is_file():
        raise ExtractionError(f"Resume file is missing on disk: {file_path.name}")
    return extract_text(file_path.read_bytes(), file_path.suffix)


def try_extract_text(content: bytes, ext: str) -> tuple[str | None, str | None]:
    """``(text, error)`` — never raises, for the upload path."""
    try:
        return extract_text(content, ext), None
    except ExtractionError as exc:
        logger.warning("Resume text extraction failed: %s", exc)
        return None, str(exc)
    except Exception as exc:  # noqa: BLE001 - defensive: never fail an upload
        logger.exception("Unexpected resume extraction failure")
        return None, f"Unexpected error while reading the document: {exc}"


__all__ = [
    "ExtractionError",
    "extract_docx_text",
    "extract_pdf_text",
    "extract_text",
    "extract_text_from_path",
    "normalize_text",
    "try_extract_text",
]
