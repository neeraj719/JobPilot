"""Render structured resume data (``ParsedResume``) to DOCX and PDF.

Rendering is separate from AI generation: it only lays out what ``parsed_data`` already
says. One neutral layout (:func:`build_layout`) feeds both writers, and the React preview
follows the same rules from the same template definition (``templates.py``).

ATS rules every output follows: one column, real headings, real bullet lists, standard
fonts, selectable text, no tables, no text boxes, no images. A test reads every export
back with ``extraction`` + ``resume_parser`` and checks the facts survive.

``python-docx`` is pinned already; ``reportlab`` is imported lazily so the app still boots
on an interpreter without it (``ExternalServiceError`` explains what is missing).
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any
from xml.sax.saxutils import escape

from app.core.exceptions import ExternalServiceError
from app.modules.ai.schemas import ParsedResume
from app.modules.resumes.recommendations import split_bullets
from app.modules.resumes.templates import SECTION_LABELS, ResumeTemplate, effective_section_order, get_template

_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_PRESENT_RE = re.compile(r"^(?:present|current|currently|now|to\s?date|till\s?date|ongoing)$", re.IGNORECASE)
_YM_RE = re.compile(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?$")
_MY_RE = re.compile(r"^(\d{1,2})/(\d{4})$")
_MONTH_NAME_RE = re.compile(r"^([A-Za-z]{3,9})\.?\s+(\d{4})$")
_MONTH_INDEX = {m.lower(): i for i, m in enumerate(_MONTHS, start=1)}

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PDF_MIME = "application/pdf"


# --------------------------------------------------------------------------- dates


def format_date(value: str | None, date_format: str = "MMM YYYY") -> str:
    """``2023-03`` -> ``Mar 2023`` / ``03/2023`` / ``2023``; anything unrecognised is kept as typed."""
    text = (value or "").strip()
    if not text:
        return ""
    if _PRESENT_RE.match(text):
        return "Present"
    year = month = None
    if m := _YM_RE.match(text):
        year, month = int(m.group(1)), int(m.group(2))
    elif m := _MY_RE.match(text):
        month, year = int(m.group(1)), int(m.group(2))
    elif m := _MONTH_NAME_RE.match(text):
        month = _MONTH_INDEX.get(m.group(1)[:3].lower())
        year = int(m.group(2))
    elif re.fullmatch(r"\d{4}", text):
        return text
    if year is None or month is None or not 1 <= month <= 12:
        return text
    if date_format == "YYYY":
        return str(year)
    if date_format == "MM/YYYY":
        return f"{month:02d}/{year}"
    return f"{_MONTHS[month - 1]} {year}"


def date_range(start: str | None, end: str | None, date_format: str) -> str:
    a, b = format_date(start, date_format), format_date(end, date_format)
    if a and b:
        return f"{a} – {b}"
    return a or b


def display_url(url: str | None) -> str:
    """``https://www.linkedin.com/in/x/`` -> ``linkedin.com/in/x`` (still recognised by the parser)."""
    text = (url or "").strip()
    text = re.sub(r"^https?://", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^www\.", "", text, flags=re.IGNORECASE)
    return text.rstrip("/")


# --------------------------------------------------------------------------- layout


@dataclass
class Entry:
    """One experience / education / project line: ``title`` bold, ``detail`` plain, ``dates`` right."""

    title: str
    detail: str = ""
    dates: str = ""
    bullets: list[str] = field(default_factory=list)
    meta: str = ""


@dataclass
class Section:
    key: str
    label: str
    paragraph: str = ""
    inline: str = ""
    bullets: list[str] = field(default_factory=list)
    entries: list[Entry] = field(default_factory=list)

    def is_empty(self) -> bool:
        return not (self.paragraph or self.inline or self.bullets or self.entries)


@dataclass
class Layout:
    name: str
    headline: str
    contacts: list[str]
    sections: list[Section]
    template: ResumeTemplate


def _clean(value: Any) -> str:
    return " ".join(str(value or "").split())


def _join(*parts: str, sep: str = " | ") -> str:
    return sep.join(p for p in (_clean(x) for x in parts) if p)


def _techs_not_in(technologies: list[str], text: str) -> list[str]:
    lowered = (text or "").lower()
    return [t for t in (_clean(x) for x in technologies or []) if t and t.lower() not in lowered]


def build_layout(parsed: ParsedResume | dict[str, Any], template: ResumeTemplate | str | None = None) -> Layout:
    """The neutral document both exporters draw (and the preview mirrors)."""
    data = parsed if isinstance(parsed, ParsedResume) else ParsedResume.model_validate(parsed or {})
    tpl = template if isinstance(template, ResumeTemplate) else get_template(template)
    fmt = tpl.date_format

    contacts = [
        c
        for c in (
            _clean(data.email),
            _clean(data.phone),
            _clean(data.location),
            display_url(data.links.linkedin),
            display_url(data.links.github),
            display_url(data.links.portfolio),
        )
        if c
    ]

    sections: list[Section] = []
    for key in effective_section_order(data.section_order, tpl):
        section = Section(key=key, label=SECTION_LABELS[key])
        if key == "summary":
            section.paragraph = _clean(data.summary)
        elif key == "skills":
            skills = [s for s in (_clean(x) for x in data.skills) if s]
            if tpl.skills_style == "bullets":
                section.bullets = skills
            else:
                section.inline = ", ".join(skills)
        elif key == "experience":
            for item in data.experience:
                if not (item.title or item.company):
                    continue
                bullets = split_bullets(item.description)
                extra = _techs_not_in(item.technologies, item.description or "")
                section.entries.append(
                    Entry(
                        title=_join(item.title, item.company),
                        detail=_clean(item.location),
                        dates=date_range(item.start, item.end, fmt),
                        bullets=bullets,
                        meta=("Technologies: " + ", ".join(extra)) if extra else "",
                    )
                )
        elif key == "projects":
            for item in data.projects:
                if not (item.name or item.description):
                    continue
                techs = _techs_not_in(item.technologies, item.description or "")
                section.entries.append(
                    Entry(
                        title=_clean(item.name) or "Project",
                        detail=display_url(item.url),
                        bullets=split_bullets(item.description),
                        meta=("Technologies: " + ", ".join(techs)) if techs else "",
                    )
                )
        elif key == "education":
            for item in data.education:
                if not (item.institution or item.degree):
                    continue
                degree = _clean(item.degree)
                field_ = _clean(item.field)
                if degree and field_ and field_.lower() not in degree.lower():
                    degree = f"{degree} in {field_}"
                section.entries.append(
                    Entry(title=_join(degree or field_, item.institution), dates=date_range(item.start, item.end, fmt))
                )
        elif key == "certifications":
            section.bullets = [
                line for line in (_join(c.name, c.issuer, c.year) for c in data.certifications if c.name) if line
            ]
        elif key == "achievements":
            section.bullets = [a for a in (_clean(x) for x in data.achievements) if a]
        elif key == "languages":
            section.inline = ", ".join(a for a in (_clean(x) for x in data.languages) if a)
        if not section.is_empty():
            sections.append(section)

    return Layout(
        name=_clean(data.name) or "Your Name",
        headline=_clean(data.headline),
        contacts=contacts,
        sections=sections,
        template=tpl,
    )


def header_contact_line(layout: Layout) -> str:
    """The line under the name. The ``two-line`` header puts the headline first on it, so
    the name stays alone on its line (parsers find the name that way)."""
    parts = list(layout.contacts)
    if layout.headline and layout.template.header_layout == "two-line":
        parts.insert(0, layout.headline)
    return " | ".join(parts)


# --------------------------------------------------------------------------- DOCX


def _hex_to_rgb(value: str) -> tuple[int, int, int]:
    text = (value or "#000000").lstrip("#")
    if len(text) != 6:
        return (0, 0, 0)
    return int(text[0:2], 16), int(text[2:4], 16), int(text[4:6], 16)


def render_docx(parsed: ParsedResume | dict[str, Any], template: ResumeTemplate | str | None = None) -> bytes:
    """A real Word document: styled paragraphs, Heading styles and List Bullet items."""
    try:
        import docx  # python-docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:  # pragma: no cover - pinned dependency
        raise ExternalServiceError("DOCX export needs python-docx (pip install python-docx).", code="dependency_missing") from exc

    layout = build_layout(parsed, template)
    tpl = layout.template
    accent = RGBColor(*_hex_to_rgb(tpl.accent))

    document = docx.Document()
    section = document.sections[0]
    margin = Inches(tpl.spacing.margin_in)
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin
    text_width = section.page_width - section.left_margin - section.right_margin

    def set_font(style_or_run, name: str) -> None:
        style_or_run.font.name = name
        element = style_or_run.element
        rpr = element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), name)

    normal = document.styles["Normal"]
    set_font(normal, tpl.fonts.docx)
    normal.font.size = Pt(tpl.sizes.body)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = tpl.spacing.line_height

    heading = document.styles["Heading 1"]
    set_font(heading, tpl.heading_fonts.docx)
    heading.font.size = Pt(tpl.sizes.heading)
    heading.font.bold = True
    heading.font.italic = False
    heading.font.color.rgb = accent if tpl.heading_style == "accent-rule" else RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(tpl.spacing.section_before)
    heading.paragraph_format.space_after = Pt(tpl.spacing.section_after)
    heading.paragraph_format.keep_with_next = True

    bullet_style = document.styles["List Bullet"]
    set_font(bullet_style, tpl.fonts.docx)
    bullet_style.font.size = Pt(tpl.sizes.body)

    def add_rule(paragraph, color: str) -> None:
        ppr = paragraph._p.get_or_add_pPr()  # noqa: SLF001
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), color.lstrip("#"))
        borders.append(bottom)
        ppr.append(borders)

    align = WD_ALIGN_PARAGRAPH.CENTER if tpl.header_layout == "centered" else WD_ALIGN_PARAGRAPH.LEFT

    name_par = document.add_paragraph()
    name_par.alignment = align
    run = name_par.add_run(layout.name)
    run.bold = True
    run.font.size = Pt(tpl.sizes.name)
    set_font(run, tpl.heading_fonts.docx)
    if layout.headline and tpl.header_layout != "two-line":
        head_par = document.add_paragraph()
        head_par.alignment = align
        head_run = head_par.add_run(layout.headline)
        head_run.font.size = Pt(tpl.sizes.headline)
    contact_line = header_contact_line(layout)
    if contact_line:
        contact_par = document.add_paragraph()
        contact_par.alignment = align
        contact_run = contact_par.add_run(contact_line)
        contact_run.font.size = Pt(tpl.sizes.small)

    for block in layout.sections:
        label = block.label.upper() if tpl.heading_style == "caps-rule" else block.label
        head = document.add_paragraph(label, style="Heading 1")
        if tpl.heading_style in ("caps-rule", "accent-rule"):
            add_rule(head, tpl.accent if tpl.heading_style == "accent-rule" else "999999")
        if block.paragraph:
            document.add_paragraph(block.paragraph)
        if block.inline:
            document.add_paragraph(block.inline)
        for item in block.bullets:
            document.add_paragraph(item, style="List Bullet")
        for position, entry in enumerate(block.entries):
            par = document.add_paragraph()
            par.paragraph_format.keep_with_next = True
            if position:
                par.paragraph_format.space_before = Pt(tpl.spacing.entry_gap)
            title = par.add_run(entry.title)
            title.bold = True
            if entry.detail:
                par.add_run(f" | {entry.detail}")
            if entry.dates:
                par.paragraph_format.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT)
                par.add_run("\t" + entry.dates)
            # The technologies line sits right under the entry line, never after the last
            # bullet, where a parser would read it as the next entry's header.
            if entry.meta:
                meta = document.add_paragraph()
                meta_run = meta.add_run(entry.meta)
                meta_run.italic = True
                meta_run.font.size = Pt(tpl.sizes.small)
            for bullet in entry.bullets:
                document.add_paragraph(bullet, style="List Bullet")

    props = document.core_properties
    props.title = f"{layout.name} - Resume"
    props.author = layout.name
    props.subject = layout.headline or "Resume"
    props.comments = ""

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- PDF


def render_pdf(parsed: ParsedResume | dict[str, Any], template: ResumeTemplate | str | None = None) -> bytes:
    """A text PDF in a standard font (selectable, extractable), laid out with reportlab platypus."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import inch
        from reportlab.pdfbase.pdfmetrics import stringWidth
        from reportlab.platypus import Flowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer
        from reportlab.platypus.flowables import HRFlowable
    except ImportError as exc:
        raise ExternalServiceError(
            "PDF export needs the reportlab package on the server (pip install reportlab==4.4.4).",
            code="dependency_missing",
        ) from exc

    layout = build_layout(parsed, template)
    tpl = layout.template
    sizes, spacing = tpl.sizes, tpl.spacing
    accent = colors.HexColor(tpl.accent)
    leading = sizes.body * spacing.line_height
    align = TA_CENTER if tpl.header_layout == "centered" else TA_LEFT

    body = ParagraphStyle("body", fontName=tpl.fonts.pdf, fontSize=sizes.body, leading=leading)
    small = ParagraphStyle("small", parent=body, fontSize=sizes.small, leading=sizes.small * spacing.line_height)
    meta_style = ParagraphStyle("meta", parent=small, fontName=tpl.fonts.pdf_italic)
    name_style = ParagraphStyle(
        "name", fontName=tpl.heading_fonts.pdf_bold, fontSize=sizes.name, leading=sizes.name * 1.2, alignment=align
    )
    headline_style = ParagraphStyle(
        "headline", parent=body, fontSize=sizes.headline, leading=sizes.headline * 1.3, alignment=align
    )
    contact_style = ParagraphStyle("contact", parent=small, alignment=align, spaceBefore=2)
    heading_style = ParagraphStyle(
        "heading",
        fontName=tpl.heading_fonts.pdf_bold,
        fontSize=sizes.heading,
        leading=sizes.heading * 1.25,
        textColor=accent if tpl.heading_style == "accent-rule" else colors.black,
        spaceBefore=spacing.section_before,
        spaceAfter=1 if tpl.heading_style != "bold" else spacing.section_after,
    )
    # The bullet glyph is part of the text (hanging indent) rather than reportlab's separate
    # bullet drawing, so text extraction reads "• item" and parsers see a list item.
    bullet_style = ParagraphStyle(
        "bullet", parent=body, leftIndent=spacing.bullet_indent, firstLineIndent=-spacing.bullet_indent * 0.75
    )

    def bullet_par(text: str):  # noqa: ANN202
        return Paragraph(f"{tpl.bullet}&nbsp;&nbsp;{escape(text)}", bullet_style)

    page_width, _ = A4
    margin = spacing.margin_in * inch
    avail = page_width - 2 * margin

    class LeftRight(Flowable):
        """Bold title (+ plain detail) on the left, dates flush right, on one baseline."""

        def __init__(self, title: str, detail: str, right: str) -> None:
            super().__init__()
            self.title, self.detail, self.right = title, detail, right

        def wrap(self, availWidth, availHeight):  # noqa: N803, ANN001, ANN201 - reportlab API
            self.width = availWidth
            self.height = leading
            return availWidth, leading

        def draw(self) -> None:
            # One text object for the whole line, so extraction reads it as one line.
            y = (leading - sizes.body) / 2 + sizes.body * 0.2
            text = self.canv.beginText(0, y)
            text.setFont(tpl.fonts.pdf_bold, sizes.body)
            text.textOut(self.title)
            if self.detail:
                text.setFont(tpl.fonts.pdf, sizes.body)
                text.textOut(f" | {self.detail}")
            if self.right:
                text.setFont(tpl.fonts.pdf, sizes.body)
                # moveCursor is relative to the start of the line: jump to the right edge.
                text.moveCursor(self.width - stringWidth(self.right, tpl.fonts.pdf, sizes.body), 0)
                text.textOut(self.right)
            self.canv.drawText(text)

    def entry_head(entry: Entry):  # noqa: ANN202
        left_width = stringWidth(entry.title, tpl.fonts.pdf_bold, sizes.body)
        if entry.detail:
            left_width += stringWidth(f" | {entry.detail}", tpl.fonts.pdf, sizes.body)
        right_width = stringWidth(entry.dates, tpl.fonts.pdf, sizes.body) if entry.dates else 0
        if left_width + right_width + 12 <= avail:
            return LeftRight(entry.title, entry.detail, entry.dates)
        # Too long for one line: one wrapping paragraph, same reading order.
        text = f"<b>{escape(entry.title)}</b>"
        if entry.detail:
            text += f" | {escape(entry.detail)}"
        if entry.dates:
            text += f" | {escape(entry.dates)}"
        return Paragraph(text, body)

    story: list[Any] = [Paragraph(escape(layout.name), name_style)]
    if layout.headline and tpl.header_layout != "two-line":
        story.append(Paragraph(escape(layout.headline), headline_style))
    contact_line = header_contact_line(layout)
    if contact_line:
        story.append(Paragraph(escape(contact_line), contact_style))

    for block in layout.sections:
        label = block.label.upper() if tpl.heading_style == "caps-rule" else block.label
        head: list[Any] = [Paragraph(escape(label), heading_style)]
        if tpl.heading_style in ("caps-rule", "accent-rule"):
            head.append(
                HRFlowable(
                    width="100%",
                    thickness=0.6,
                    color=accent if tpl.heading_style == "accent-rule" else colors.HexColor("#9ca3af"),
                    spaceBefore=0,
                    spaceAfter=spacing.section_after,
                )
            )
        content: list[Any] = []
        if block.paragraph:
            content.append(Paragraph(escape(block.paragraph), body))
        if block.inline:
            content.append(Paragraph(escape(block.inline), body))
        for item in block.bullets:
            content.append(bullet_par(item))
        for position, entry in enumerate(block.entries):
            parts: list[Any] = []
            if position:
                parts.append(Spacer(1, spacing.entry_gap))
            parts.append(entry_head(entry))
            if entry.meta:
                parts.append(Paragraph(escape(entry.meta), meta_style))
            for bullet in entry.bullets:
                parts.append(bullet_par(bullet))
            content.append(KeepTogether(parts) if len(parts) <= 6 else parts[0])
            if len(parts) > 6:
                content.extend(parts[1:])
        # Keep a heading with the start of its section.
        if content:
            story.append(KeepTogether([*head, content[0]]))
            story.extend(content[1:])
        else:
            story.extend(head)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title=f"{layout.name} - Resume",
        author=layout.name,
        subject=layout.headline or "Resume",
        creator="JobPilot AI",
    )
    doc.build(story)
    return buffer.getvalue()


def render(parsed: ParsedResume | dict[str, Any], fmt: str, template: ResumeTemplate | str | None = None) -> tuple[bytes, str]:
    """``(bytes, mime type)`` for ``fmt`` = ``pdf`` | ``docx``."""
    kind = (fmt or "").lower()
    if kind == "pdf":
        return render_pdf(parsed, template), PDF_MIME
    if kind == "docx":
        return render_docx(parsed, template), DOCX_MIME
    raise ValueError(f"Unsupported export format: {fmt!r}")


__all__ = [
    "DOCX_MIME",
    "PDF_MIME",
    "Entry",
    "Layout",
    "Section",
    "build_layout",
    "date_range",
    "display_url",
    "format_date",
    "render",
    "render_docx",
    "render_pdf",
]
